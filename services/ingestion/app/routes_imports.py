import csv
import httpx
import polars as pl
from fastapi import File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from io import BytesIO, StringIO
from openpyxl import Workbook
from starlette.concurrency import run_in_threadpool
from typing import Any
from urllib.parse import quote
from fastapi import APIRouter

from . import import_store, supabase_client, table_reading
from .import_store import safe_workbook_sheet_name, spreadsheet_export_value
from .models import LinkRequest
from .remote_sources import discover_page_candidates, fetch_link_source, is_page_response, parse_drive_url
from .source_cache import persist_cached_source, profile_cached_source, remember_source
from .source_limits import read_upload
from .supabase_client import upstream_error
from .table_reading import recommend_indicators, suggest_mapping

router = APIRouter()


@router.get("/imports")
def list_imports(status: str | None = None, limit: int = 100) -> list[dict[str, Any]]:
    query_parameters = {"select": "id,title,file_name,source_url,status,total_rows,created_at,updated_at", "order": "created_at.desc", "limit": str(max(1, min(limit, 500))), "status": "neq.discarded"}
    if status in {"draft", "analyzing", "needs_review", "approved", "archived"}:
        query_parameters["status"] = f"eq.{status}"
    imports_response = httpx.get(supabase_client.supabase_url("/rest/v1/imports"), params=query_parameters, headers=supabase_client.supabase_headers(), timeout=30.0)
    if not imports_response.is_success:
        raise upstream_error("Não foi possível carregar as importações salvas.", imports_response)
    return imports_response.json()

@router.post("/profile")
async def profile_file(file: UploadFile | None = File(None), sheet_name: str | None = Form(None), upload_token: str | None = Form(None)) -> dict[str, Any]:
    if upload_token:
        return await run_in_threadpool(profile_cached_source, upload_token, "uploaded_file", sheet_name)
    if file is None or not file.filename:
        raise HTTPException(400, "Envie um arquivo para analisar.")
    source_content = await read_upload(file)
    upload_token = remember_source(file.filename, source_content, None)
    return await run_in_threadpool(profile_cached_source, upload_token, "uploaded_file", sheet_name)

@router.post("/profile-link")
async def profile_link(link_request: LinkRequest) -> dict[str, Any]:
    if link_request.upload_token:
        return await run_in_threadpool(profile_cached_source, link_request.upload_token, "remote_file", link_request.sheet_name)
    source_name, source_content, content_type, final_url, folder_candidates = await fetch_link_source(str(link_request.source_url))
    if folder_candidates is not None:
        return {"kind": "web_page", "source_url": final_url, "download_candidates": folder_candidates}
    if not parse_drive_url(str(link_request.source_url)) and is_page_response(source_name, content_type):
        page_candidates = await run_in_threadpool(discover_page_candidates, source_content, final_url)
        if page_candidates:
            return {"kind": "web_page", "source_url": final_url, "download_candidates": page_candidates}
        raise HTTPException(415, "A página não apresentou um arquivo CSV, XLSX, XLS, JSON ou ZIP identificável.")
    upload_token = remember_source(source_name, source_content, final_url)
    return await run_in_threadpool(profile_cached_source, upload_token, "remote_file", link_request.sheet_name)

@router.post("/imports/draft")
async def create_import_draft(file: UploadFile | None = File(None), dataset_id: str | None = Form(None), title: str | None = Form(None), reference_year: int | None = Form(None), sheet_name: str | None = Form(None), include_all_sheets: bool = Form(False), upload_token: str | None = Form(None)) -> dict[str, Any]:
    if upload_token:
        return await run_in_threadpool(persist_cached_source, upload_token, dataset_id, title, reference_year, sheet_name, include_all_sheets)
    if file is None or not file.filename:
        raise HTTPException(400, "Envie um arquivo para criar o rascunho.")
    source_content = await read_upload(file)
    return await run_in_threadpool(import_store.persist_import, file.filename, source_content, None, dataset_id, title or file.filename, reference_year, sheet_name, include_all_sheets)

@router.post("/imports/draft-link")
async def create_link_import_draft(link_request: LinkRequest) -> dict[str, Any]:
    if link_request.upload_token:
        return await run_in_threadpool(persist_cached_source, link_request.upload_token, link_request.dataset_id, link_request.title, link_request.reference_year, link_request.sheet_name, link_request.include_all_sheets)
    drive_target = parse_drive_url(str(link_request.source_url))
    if drive_target and drive_target[0] == "folder":
        raise HTTPException(415, "Escolha um arquivo específico da pasta do Google Drive antes de criar o rascunho.")
    source_name, source_content, content_type, final_url, _ = await fetch_link_source(str(link_request.source_url))
    if not drive_target and is_page_response(source_name, content_type):
        raise HTTPException(415, "Use o link direto de um arquivo para criar o rascunho.")
    return await run_in_threadpool(import_store.persist_import, source_name, source_content, final_url, link_request.dataset_id, link_request.title or source_name, link_request.reference_year, link_request.sheet_name, link_request.include_all_sheets)

@router.post("/imports/{import_id}/discard")
def discard_import(import_id: str) -> dict[str, str]:
    import_response = httpx.get(supabase_client.supabase_url("/rest/v1/imports"), params={"id": f"eq.{import_id}", "select": "storage_path"}, headers=supabase_client.supabase_headers(), timeout=30.0)
    if not import_response.is_success or not import_response.json():
        raise HTTPException(404, "Importação não encontrada.")
    storage_path = import_response.json()[0].get("storage_path")
    if storage_path:
        storage_response = httpx.delete(supabase_client.supabase_url(f"/storage/v1/object/source-files/{storage_path}"), headers=supabase_client.supabase_headers(), timeout=30.0)
        if not storage_response.is_success:
            raise HTTPException(502, "Não foi possível remover o arquivo privado da importação.")
    discard_response = httpx.patch(supabase_client.supabase_url(f"/rest/v1/imports?id=eq.{import_id}"), headers=supabase_client.supabase_headers(), json={"status": "discarded", "storage_path": None}, timeout=30.0)
    if not discard_response.is_success:
        raise HTTPException(502, "Não foi possível descartar a importação.")
    return {"import_id": import_id, "status": "discarded"}

@router.get("/imports/{import_id}/validation-issues")
def list_import_validation_issues(import_id: str) -> list[dict[str, Any]]:
    validation_issues: list[dict[str, Any]] = []
    page_size = 1000
    page_offset = 0
    while True:
        issues_response = httpx.get(
            supabase_client.supabase_url("/rest/v1/validation_issues"),
            params={"import_id": f"eq.{import_id}", "select": "severity,row_number,field,message,created_at", "order": "row_number.asc,created_at.desc", "limit": str(page_size), "offset": str(page_offset)},
            headers=supabase_client.supabase_headers(),
            timeout=30.0,
        )
        if not issues_response.is_success:
            raise HTTPException(502, "Não foi possível carregar as pendências desta importação.")
        current_page = issues_response.json()
        validation_issues.extend(current_page)
        if len(current_page) < page_size:
            return validation_issues
        page_offset += page_size

@router.get("/imports/{import_id}/sheets")
def list_import_sheets(import_id: str, sheet_name: str | None = None) -> dict[str, Any]:
    sheets_response = httpx.get(
        supabase_client.supabase_url("/rest/v1/import_sheets"),
        params={"import_id": f"eq.{import_id}", "select": "sheet_name,sheet_position,row_count,column_count,columns_profile,sample_rows,selected_for_treatment", "order": "sheet_position.asc"},
        headers=supabase_client.supabase_headers(),
        timeout=30.0,
    )
    if not sheets_response.is_success:
        raise HTTPException(502, "Não foi possível carregar as abas desta importação.")
    sheet_records = sheets_response.json()
    selected_record = next((record for record in sheet_records if record["sheet_name"] == sheet_name), None)
    if sheet_name and not selected_record:
        raise HTTPException(404, "A aba escolhida não faz parte desta importação.")
    profile = None
    if selected_record:
        stored_columns = selected_record.get("columns_profile") or []
        stored_sample = selected_record.get("sample_rows") or []
        if stored_columns and stored_sample:
            profile = {"columns": stored_columns, "sample": stored_sample}
        else:
            rows_response = httpx.get(
                supabase_client.supabase_url("/rest/v1/import_rows"),
                params={"import_id": f"eq.{import_id}", "sheet_name": f"eq.{sheet_name}", "select": "raw_row", "order": "row_number.asc", "limit": "20"},
                headers=supabase_client.supabase_headers(),
                timeout=30.0,
            )
            if not rows_response.is_success:
                raise HTTPException(502, "Não foi possível carregar uma prévia da aba escolhida.")
            sample_rows = [record.get("raw_row") or {} for record in rows_response.json()]
            if sample_rows:
                sample_table = pl.from_dicts(sample_rows, infer_schema_length=20, strict=False)
                profile = {
                    "columns": [{"name": column_name, "dtype": str(sample_table[column_name].dtype), "null_count": int(sample_table[column_name].null_count())} for column_name in sample_table.columns],
                    "sample": sample_rows,
                    "suggestions": suggest_mapping(sample_table),
                    "indicator_recommendations": recommend_indicators(sample_table),
                }
    return {
        "sheets": [{"name": record["sheet_name"], "rows": record["row_count"], "columns": record["column_count"], "has_data": record["row_count"] > 0, "imported": record["row_count"] > 0} for record in sheet_records],
        "profile": profile,
    }

@router.get("/imports/{import_id}/resume")
def resume_import(import_id: str) -> dict[str, Any]:
    import_response = httpx.get(
        supabase_client.supabase_url("/rest/v1/imports"),
        params={"id": f"eq.{import_id}", "select": "id,file_name,source_url,profile,status,total_rows"},
        headers=supabase_client.supabase_headers(),
        timeout=30.0,
    )
    if not import_response.is_success or not import_response.json():
        raise HTTPException(404, "Não encontramos essa importação salva.")
    import_record = import_response.json()[0]
    saved_profile = import_record.get("profile") or {}
    selected_sheet = saved_profile.get("selected_sheet")
    sheet_data = list_import_sheets(import_id, selected_sheet)
    resumed_profile = {**saved_profile, "sheets": sheet_data["sheets"]}
    if selected_sheet and sheet_data["profile"]:
        resumed_profile.update(sheet_data["profile"])
    return {"import_id": import_id, "status": import_record["status"], "profile": resumed_profile}

@router.post("/imports/{import_id}/include-sheets")
def include_remaining_import_sheets(import_id: str) -> dict[str, Any]:
    import_response = httpx.get(
        supabase_client.supabase_url("/rest/v1/imports"),
        params={"id": f"eq.{import_id}", "select": "storage_path,file_name,profile,total_rows"},
        headers=supabase_client.supabase_headers(),
        timeout=30.0,
    )
    if not import_response.is_success or not import_response.json():
        raise HTTPException(404, "Importação não encontrada.")
    import_record = import_response.json()[0]
    storage_path = import_record.get("storage_path")
    if not storage_path:
        raise HTTPException(422, "O arquivo original não está disponível para importar outras abas.")
    file_response = httpx.get(
        supabase_client.supabase_url(f"/storage/v1/object/source-files/{quote(storage_path, safe='/')}"),
        headers=supabase_client.supabase_headers(),
        timeout=180.0,
    )
    if not file_response.is_success:
        raise HTTPException(502, "Não foi possível abrir o arquivo original para importar as outras abas.")
    existing_response = httpx.get(
        supabase_client.supabase_url("/rest/v1/import_sheets"),
        params={"import_id": f"eq.{import_id}", "select": "sheet_name,row_count"},
        headers=supabase_client.supabase_headers(),
        timeout=30.0,
    )
    if not existing_response.is_success:
        raise HTTPException(502, "Não foi possível verificar quais abas já foram importadas.")
    existing_rows = {record["sheet_name"]: record["row_count"] for record in existing_response.json()}
    file_name = import_record["file_name"]
    original_profile = import_record.get("profile") or {}
    selected_sheet = original_profile.get("selected_sheet")
    source_tables = table_reading.import_tables(file_name, file_response.content, selected_sheet, True)
    pending_tables = [(name, table) for name, table in source_tables if not existing_rows.get(name)]
    staged_rows = [
        {"import_id": import_id, "sheet_name": table_name, "row_number": row_number, "raw_row": row_values}
        for table_name, table in pending_tables
        for row_number, row_values in enumerate(table.to_dicts(), start=1)
    ]
    for start_index in range(0, len(staged_rows), 500):
        rows_response = httpx.post(
            supabase_client.supabase_url("/rest/v1/import_rows"),
            params={"on_conflict": "import_id,sheet_name,row_number"},
            headers=supabase_client.supabase_headers("resolution=ignore-duplicates"),
            json=staged_rows[start_index:start_index + 500],
            timeout=60.0,
        )
        if not rows_response.is_success:
            raise HTTPException(502, "Não foi possível guardar as linhas das abas adicionais.")
    sheet_records = []
    for sheet_position, (table_name, table) in enumerate(source_tables, start=1):
        sheet_records.append({
            "import_id": import_id,
            "sheet_name": table_name,
            "sheet_position": sheet_position,
            "row_count": table.height,
            "column_count": table.width,
            "columns_profile": [{"name": column_name, "dtype": str(table[column_name].dtype), "null_count": int(table[column_name].null_count())} for column_name in table.columns],
            "sample_rows": table.head(20).to_dicts(),
            "selected_for_treatment": False,
        })
    sheets_response = httpx.post(
        supabase_client.supabase_url("/rest/v1/import_sheets"),
        params={"on_conflict": "import_id,sheet_name"},
        headers=supabase_client.supabase_headers("resolution=merge-duplicates"),
        json=sheet_records,
        timeout=30.0,
    )
    if not sheets_response.is_success:
        raise HTTPException(502, "As linhas foram preparadas, mas não foi possível atualizar a lista de abas.")
    additional_rows = sum(table.height for _, table in pending_tables)
    if additional_rows:
        total_rows = int(import_record.get("total_rows") or 0) + additional_rows
        total_response = httpx.patch(
            supabase_client.supabase_url(f"/rest/v1/imports?id=eq.{import_id}"),
            headers=supabase_client.supabase_headers(),
            json={"total_rows": total_rows},
            timeout=30.0,
        )
        if not total_response.is_success:
            raise HTTPException(502, "As abas foram guardadas, mas não foi possível atualizar o total de linhas.")
    return {"import_id": import_id, "added_rows": additional_rows, "added_sheets": [name for name, _ in pending_tables]}

@router.get("/imports/{import_id}/export.csv")
def export_import_csv(import_id: str) -> Response:
    import_rows = import_store.fetch_import_rows(import_id)
    field_names = list(dict.fromkeys(field_name for import_row in import_rows for field_name in import_row["raw_row"]))
    sheet_field_name = "aba"
    suffix = 2
    while sheet_field_name in set(field_names):
        sheet_field_name = f"aba_{suffix}"
        suffix += 1
    csv_buffer = StringIO()
    csv_writer = csv.DictWriter(csv_buffer, fieldnames=[sheet_field_name, *field_names], extrasaction="ignore", delimiter=";")
    csv_writer.writeheader()
    for import_row in import_rows:
        csv_writer.writerow({sheet_field_name: import_row["sheet_name"], **{field_name: spreadsheet_export_value(field_value) for field_name, field_value in import_row["raw_row"].items()}})
    return Response("\ufeff" + csv_buffer.getvalue(), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="importacao-{import_id}.csv"'})

@router.get("/imports/{import_id}/export.xlsx")
def export_import_xlsx(import_id: str) -> Response:
    import_rows = import_store.fetch_import_rows(import_id)
    workbook = Workbook()
    workbook.remove(workbook.active)
    import_rows_by_sheet: dict[str, list[dict[str, Any]]] = {}
    for import_row in import_rows:
        import_rows_by_sheet.setdefault(import_row["sheet_name"], []).append(import_row)
    if not import_rows_by_sheet:
        import_rows_by_sheet["Dados"] = []
    for sheet_name, sheet_rows in import_rows_by_sheet.items():
        worksheet = workbook.create_sheet(safe_workbook_sheet_name(sheet_name, workbook.sheetnames))
        records = [import_row["raw_row"] for import_row in sheet_rows]
        field_names = list(dict.fromkeys(field_name for record in records for field_name in record))
        worksheet.append(field_names)
        for record in records:
            worksheet.append([spreadsheet_export_value(record.get(field_name)) for field_name in field_names])
        worksheet.freeze_panes = "A2"
        for column_cells in worksheet.columns:
            longest_value = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(longest_value + 2, 48)
    workbook_buffer = BytesIO()
    workbook.save(workbook_buffer)
    return Response(workbook_buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="importacao-{import_id}.xlsx"'})
