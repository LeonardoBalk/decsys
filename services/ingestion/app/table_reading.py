import csv
import json
import math
import polars as pl
import re
import unicodedata
import zipfile
from datetime import date, datetime, time as datetime_time
from decimal import Decimal
from fastapi import HTTPException
from io import BytesIO, StringIO
from itertools import islice
from openpyxl import load_workbook
from pathlib import Path
from typing import Any

from . import ai_assessment
from .config import ai_assessment_size_limit, supported_extensions
from .source_limits import bounded_gzip_decompress, ensure_source_size


def import_tables(source_name: str, source_content: bytes, selected_sheet: str | None, include_all_sheets: bool) -> list[tuple[str, pl.DataFrame]]:
    if not include_all_sheets or not selected_sheet:
        selected_sheet = resolve_sheet_name(source_name, source_content, selected_sheet)
    if include_all_sheets and Path(source_name).suffix.lower() == ".xlsx":
        tables: list[tuple[str, pl.DataFrame]] = []
        for sheet in workbook_sheets(source_content):
            try:
                source_table = normalize_table_columns(read_table(source_name, source_content, str(sheet["name"])))
            except HTTPException:
                continue
            if source_table.height and source_table.width:
                tables.append((str(sheet["name"]), source_table))
        if tables:
            return tables
        raise HTTPException(422, "Nenhuma aba da planilha contém uma tabela que possa ser importada.")
    source_table = normalize_table_columns(read_table(source_name, source_content, selected_sheet))
    return [(selected_sheet or "Dados", source_table)]

def json_compatible_value(field_value: Any) -> Any:
    if isinstance(field_value, (datetime, date, datetime_time)):
        return field_value.isoformat()
    if isinstance(field_value, Decimal):
        return format(field_value, "f")
    if isinstance(field_value, float) and not math.isfinite(field_value):
        return None
    if isinstance(field_value, dict):
        return {str(field_name): json_compatible_value(field_value) for field_name, field_value in field_value.items()}
    if isinstance(field_value, (list, tuple)):
        return [json_compatible_value(value) for value in field_value]
    return field_value

def normalize_column(column_name: str) -> str:
    normalized_name = "".join(character for character in unicodedata.normalize("NFKD", column_name.strip().lower()) if not unicodedata.combining(character))
    normalized_name = re.sub(r"_duplicated_(\d+)$", lambda match: f"_{int(match.group(1)) + 2}", normalized_name)
    return "_".join(normalized_name.replace("/", " ").split())

def normalize_table_columns(source_table: pl.DataFrame) -> pl.DataFrame:
    normalized_names: list[str] = []
    used_names: set[str] = set()
    for position, column_name in enumerate(source_table.columns, start=1):
        normalized_name = normalize_column(column_name)
        if not normalized_name or normalized_name.startswith("__unnamed__"):
            normalized_name = f"coluna_{position}"
        unique_name = normalized_name
        suffix = 2
        while unique_name in used_names:
            unique_name = f"{normalized_name}_{suffix}"
            suffix += 1
        normalized_names.append(unique_name)
        used_names.add(unique_name)
    normalized_table = source_table.rename(dict(zip(source_table.columns, normalized_names)))
    for column_name in [column_name for column_name in normalized_table.columns if is_municipality_code_column(column_name)]:
        code_as_text = pl.col(column_name).cast(pl.String)
        normalized_table = normalized_table.with_columns(
            pl.when(code_as_text.str.contains(r"^[0-9]{6,7}\.0+$"))
            .then(code_as_text.str.replace(r"\.0+$", ""))
            .otherwise(code_as_text)
            .alias(column_name)
        )
    return normalized_table

def is_municipality_code_column(column_name: str) -> bool:
    return re.search(r"ibge|cod(?:igo)?_(?:do_|de_)?mun|municipio_cod", column_name) is not None

def period_from_column_name(column_name: str) -> tuple[int, int] | None:
    month_numbers = {"janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6, "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12}
    normalized_name = normalized_sheet_title(normalize_column(column_name))
    period_match = re.search(r"(" + "|".join(month_numbers) + r")_(20[0-9]{2})", normalized_name)
    if not period_match:
        return None
    return int(period_match.group(2)), month_numbers[period_match.group(1)]

def decode_csv_text(source_content: bytes) -> str:
    for candidate_encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return source_content.decode(candidate_encoding)
        except UnicodeDecodeError:
            continue
    return source_content.decode("utf-8", errors="replace")

def detect_csv_separator(csv_text: str) -> str:
    sample_text = csv_text[:8192]
    try:
        return csv.Sniffer().sniff(sample_text, delimiters=";,\t|").delimiter
    except csv.Error:
        candidate_counts = {candidate: sample_text.count(candidate) for candidate in (";", ",", "\t", "|")}
        return max(candidate_counts, key=candidate_counts.get) if any(candidate_counts.values()) else ","

def has_wrapped_csv_records(csv_text: str, separator: str) -> bool:
    sample_lines = list(islice(StringIO(csv_text), 7))
    if len(sample_lines) < 4:
        return False
    header_fields = next(csv.reader([sample_lines[0]], delimiter=separator), [])
    if len(header_fields) < 2:
        return False
    for source_line in sample_lines[1:]:
        source_record = source_line.rstrip("\r\n")
        parsed_record = next(csv.reader([source_record], delimiter=separator), [])
        if len(parsed_record) != 1 or not source_record.startswith('"') or not source_record.endswith('"'):
            return False
        unwrapped_record = source_record[1:-1].replace('""', '"')
        if len(next(csv.reader([unwrapped_record], delimiter=separator), [])) != len(header_fields):
            return False
    return True

def unwrap_csv_records(csv_text: str) -> str:
    normalized_lines = StringIO()
    for source_line in StringIO(csv_text):
        source_record = source_line.rstrip("\r\n")
        line_ending = source_line[len(source_record):]
        if source_record.startswith('"') and source_record.endswith('"'):
            source_record = source_record[1:-1].replace('""', '"')
        normalized_lines.write(source_record + line_ending)
    return normalized_lines.getvalue()

def read_csv_source_table(csv_text: str, separator: str) -> pl.DataFrame:
    try:
        return pl.read_csv(StringIO(csv_text), try_parse_dates=True, infer_schema_length=500, separator=separator)
    except pl.exceptions.ComputeError:
        return pl.read_csv(StringIO(csv_text), try_parse_dates=True, infer_schema_length=None, separator=separator)

def normalize_brazilian_decimals(source_table: pl.DataFrame) -> pl.DataFrame:
    numeric_pattern = r"^-?(?:(?:[0-9]{1,3}(\.[0-9]{3})*|[0-9]+)(,[0-9]+)?|[0-9]+\.[0-9]+|,[0-9]+)$"
    normalized_columns: list[pl.Expr] = []
    for column_name, data_type in zip(source_table.columns, source_table.dtypes):
        if data_type != pl.String:
            continue
        non_empty_values = [str(value).strip() for value in source_table[column_name].drop_nulls().to_list() if str(value).strip() not in {"", "-"}]
        numeric_values = [value for value in non_empty_values if re.fullmatch(numeric_pattern, value)]
        has_decimal_separator = any("," in value or "." in value for value in numeric_values)
        if non_empty_values and has_decimal_separator and len(numeric_values) / len(non_empty_values) >= 0.6:
            trimmed_values = pl.col(column_name).cast(pl.String).str.strip_chars()
            brazilian_decimal_values = trimmed_values.str.replace_all(".", "", literal=True).str.replace(",", ".", literal=True)
            normalized_columns.append(
                pl.when(trimmed_values.is_in(["", "-"]))
                .then(None)
                .when(trimmed_values.str.contains(","))
                .then(brazilian_decimal_values)
                .otherwise(trimmed_values)
                .cast(pl.Float64, strict=False)
                .alias(column_name)
            )
    return source_table.with_columns(normalized_columns) if normalized_columns else source_table

def read_json_table(source_content: bytes) -> pl.DataFrame:
    source_records = json.loads(source_content)
    if isinstance(source_records, list) and source_records and isinstance(source_records[0], dict):
        label_record = source_records[0]
        has_sidra_labels = "V" in label_record and any(str(value).endswith("(Código)") for value in label_record.values())
        if has_sidra_labels:
            source_table = pl.from_dicts(source_records[1:])
            readable_names = unique_header_names([str(label_record.get(column_name, column_name)) for column_name in source_table.columns])
            return source_table.rename(dict(zip(source_table.columns, readable_names)))
    return pl.read_json(BytesIO(source_content))

def normalized_sheet_title(sheet_title: str) -> str:
    return "".join(character for character in unicodedata.normalize("NFKD", sheet_title).lower() if not unicodedata.combining(character)).strip()

def is_navigation_sheet(sheet_title: str) -> bool:
    return normalized_sheet_title(sheet_title) in {"sumario", "indice", "contents"}

def workbook_sheets(source_content: bytes) -> list[dict[str, int | str]]:
    workbook = load_workbook(BytesIO(source_content), read_only=True, data_only=True)
    sheet_profiles = []
    for worksheet in workbook.worksheets:
        if is_navigation_sheet(worksheet.title):
            continue
        row_count = worksheet.max_row
        column_count = worksheet.max_column
        if row_count is None or column_count is None:
            row_count = 0
            column_count = 0
            for row_number, row_values in enumerate(worksheet.iter_rows(values_only=True), start=1):
                populated_columns = [column_number for column_number, cell_value in enumerate(row_values, start=1) if has_cell_value(cell_value)]
                if populated_columns:
                    row_count = row_number
                    column_count = max(column_count, max(populated_columns))
            if row_count and column_count:
                sheet_profiles.append({"name": worksheet.title, "rows": row_count, "columns": column_count})
            continue
        if any(any(has_cell_value(cell_value) for cell_value in row_values) for row_values in worksheet.iter_rows(values_only=True)):
            sheet_profiles.append({"name": worksheet.title, "rows": row_count, "columns": column_count})
    return sheet_profiles

def has_cell_value(cell_value: Any) -> bool:
    return cell_value is not None and str(cell_value).strip() != ""

def has_excel_cell_value(cell_value: Any) -> bool:
    return has_cell_value(cell_value) and not str(cell_value).strip().lower().startswith("__unnamed__")

def is_numeric_cell(cell_value: Any) -> bool:
    if isinstance(cell_value, (int, float)) and not isinstance(cell_value, bool):
        return True
    if not isinstance(cell_value, str):
        return False
    numeric_text = cell_value.strip().replace(" ", "")
    if "," in numeric_text and "." in numeric_text:
        if numeric_text.rfind(",") > numeric_text.rfind("."):
            numeric_text = numeric_text.replace(".", "").replace(",", ".")
        else:
            numeric_text = numeric_text.replace(",", "")
    else:
        numeric_text = numeric_text.replace(",", ".")
    return re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", numeric_text) is not None

def unique_header_names(header_names: list[str]) -> list[str]:
    resolved_names: list[str] = []
    used_names: set[str] = set()
    for position, header_name in enumerate(header_names, start=1):
        base_name = header_name or f"Coluna {position}"
        unique_name = base_name
        suffix = 2
        while unique_name in used_names:
            unique_name = f"{base_name} ({suffix})"
            suffix += 1
        resolved_names.append(unique_name)
        used_names.add(unique_name)
    return resolved_names

def infer_excel_headers(source_content: bytes, sheet_name: str) -> dict[str, Any] | None:
    workbook = load_workbook(BytesIO(source_content), read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    sample_rows = list(islice(worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row or 40, 40), values_only=True), 40))
    sample_column_count = max((len(row_values) for row_values in sample_rows), default=0)
    data_row_number = None
    for row_number, row_values in enumerate(sample_rows, start=1):
        populated_values = [cell_value for cell_value in row_values if has_excel_cell_value(cell_value)]
        numeric_count = sum(is_numeric_cell(cell_value) for cell_value in populated_values)
        if len(populated_values) >= 3 and numeric_count >= 2 and numeric_count / len(populated_values) >= 0.2:
            data_row_number = row_number
            break
    if data_row_number is None:
        populated_counts = [sum(has_excel_cell_value(cell_value) for cell_value in row_values) for row_values in sample_rows]
        maximum_populated_count = max(populated_counts, default=0)
        if maximum_populated_count < 2:
            return None
        header_row_index = populated_counts.index(maximum_populated_count)
        header_row = sample_rows[header_row_index]
        header_names = [" ".join(str(cell_value).split()) if has_cell_value(cell_value) else "" for cell_value in header_row]
        data_samples = sample_rows[header_row_index + 1:min(len(sample_rows), header_row_index + 6)]
        has_following_data = any(sum(has_excel_cell_value(cell_value) for cell_value in row_values) >= 2 for row_values in data_samples)
        if not has_following_data:
            return None
        data_columns = [column_index for column_index in range(sample_column_count) if any(column_index < len(row_values) and has_excel_cell_value(row_values[column_index]) for row_values in data_samples)]
        resolved_headers = [header_names[column_index] if column_index < len(header_names) else "" for column_index in data_columns]
        return {"header_row": header_row_index, "header_rows": [header_row_index + 1], "header_names": unique_header_names(resolved_headers)}
    if data_row_number == 1:
        return None
    header_start = data_row_number - 1
    while header_start > 1 and data_row_number - header_start < 3:
        previous_header_row = sample_rows[header_start - 2]
        populated_header_cells = sum(has_excel_cell_value(cell_value) for cell_value in previous_header_row)
        if populated_header_cells < 2:
            break
        header_start -= 1
    header_rows = sample_rows[header_start - 1:data_row_number - 1]
    data_samples = sample_rows[data_row_number - 1:min(len(sample_rows), data_row_number + 5)]
    carried_labels = ["" for _ in header_rows]
    header_names: list[str] = []
    for column_index in range(sample_column_count):
        labels: list[str] = []
        has_direct_header = False
        for level, header_row in enumerate(header_rows):
            cell_value = header_row[column_index] if column_index < len(header_row) else None
            if has_excel_cell_value(cell_value):
                has_direct_header = True
                header_text = str(cell_value).replace("-_x000d_\n", "").replace("-\n", "").replace("_x000d_", " ").replace("\r", " ").replace("\n", " ")
                carried_labels[level] = " ".join(header_text.split())
            if carried_labels[level] and carried_labels[level] not in labels:
                labels.append(carried_labels[level])
        has_data = any(column_index < len(data_row) and has_excel_cell_value(data_row[column_index]) for data_row in data_samples)
        if has_data:
            header_names.append(" ".join(labels) if has_direct_header else "")
    return {"header_row": data_row_number - 2, "header_rows": list(range(header_start, data_row_number)), "header_names": unique_header_names(header_names)}

def trim_trailing_spreadsheet_notes(source_table: pl.DataFrame) -> pl.DataFrame:
    if not source_table.height or not source_table.width:
        return source_table
    populated_fields = source_table.select(pl.sum_horizontal(*[pl.col(column_name).is_not_null().cast(pl.UInt16) for column_name in source_table.columns]).alias("populated_fields")).get_column("populated_fields").to_list()
    minimum_populated_fields = min(2, source_table.width)
    last_data_position = next((position for position in range(len(populated_fields) - 1, -1, -1) if populated_fields[position] >= minimum_populated_fields), -1)
    return source_table.slice(0, last_data_position + 1)

def read_excel_table(source_content: bytes, sheet_name: str) -> pl.DataFrame:
    header_configuration = infer_excel_headers(source_content, sheet_name)
    read_options = {"header_row": header_configuration["header_row"]} if header_configuration else {}
    try:
        source_table = pl.read_excel(BytesIO(source_content), sheet_name=sheet_name, infer_schema_length=500, read_options=read_options)
    except pl.exceptions.NoDataError:
        raise HTTPException(422, f'A aba "{sheet_name}" não possui dados para importar. Escolha outra aba da planilha.')
    if header_configuration and len(header_configuration["header_names"]) == source_table.width:
        source_table = source_table.rename(dict(zip(source_table.columns, header_configuration["header_names"])))
        repeated_header_row = pl.all_horizontal([
            pl.col(column_name).cast(pl.String).str.strip_chars().str.to_lowercase() == header_name.strip().lower()
            for column_name, header_name in zip(source_table.columns, header_configuration["header_names"])
        ])
        source_table = source_table.filter(~repeated_header_row)
    return trim_trailing_spreadsheet_notes(source_table)

def resolve_sheet_name(source_name: str, source_content: bytes, requested_sheet_name: str | None) -> str | None:
    source_extension = Path(source_name).suffix.lower()
    if source_extension not in {".xlsx", ".xls"}:
        return None
    if source_extension == ".xls":
        return requested_sheet_name
    available_sheets = workbook_sheets(source_content)
    if requested_sheet_name:
        if requested_sheet_name not in {str(sheet["name"]) for sheet in available_sheets}:
            raise HTTPException(422, "A aba selecionada não existe mais nessa planilha.")
        return requested_sheet_name
    sheet_candidates: list[tuple[int, str]] = []
    for sheet_candidate in available_sheets:
        try:
            candidate_table = read_table(source_name, source_content, str(sheet_candidate["name"]))
        except HTTPException:
            continue
        if candidate_table.height and candidate_table.width:
            sheet_candidates.append((candidate_table.height * candidate_table.width, str(sheet_candidate["name"])))
    if sheet_candidates:
        return max(sheet_candidates, key=lambda candidate: candidate[0])[1]
    raise HTTPException(422, "Nenhuma aba da planilha contém uma tabela que possa ser lida.")

def read_table(source_name: str, source_content: bytes, sheet_name: str | None = None) -> pl.DataFrame:
    if not source_content:
        raise HTTPException(422, "O arquivo está vazio. Envie uma planilha ou arquivo de dados válido.")
    try:
        return read_table_content(source_name, source_content, sheet_name)
    except HTTPException:
        raise
    except Exception as parse_error:
        raise HTTPException(422, "Não foi possível ler este arquivo. Confirme se ele não está corrompido e se o formato corresponde à extensão.") from parse_error

def read_table_content(source_name: str, source_content: bytes, sheet_name: str | None = None) -> pl.DataFrame:
    source_extension = Path(source_name).suffix.lower()
    if source_extension == ".gz":
        return read_table(Path(source_name).stem, bounded_gzip_decompress(source_content))
    if source_extension == ".zip":
        with zipfile.ZipFile(BytesIO(source_content)) as source_archive:
            candidate_members = [member for member in source_archive.namelist() if not member.endswith("/") and Path(member).suffix.lower() in supported_extensions]
            if not candidate_members:
                raise HTTPException(415, "O arquivo .zip não contém nenhum CSV, XLSX, XLS ou JSON.")
            chosen_member = candidate_members[0]
            ensure_source_size(source_archive.getinfo(chosen_member).file_size)
            return read_table(Path(chosen_member).name, source_archive.read(chosen_member))
    if source_extension == ".csv":
        csv_text = decode_csv_text(source_content)
        separator = detect_csv_separator(csv_text)
        if has_wrapped_csv_records(csv_text, separator):
            csv_text = unwrap_csv_records(csv_text)
        return normalize_brazilian_decimals(read_csv_source_table(csv_text, separator))
    if source_extension == ".xlsx":
        selected_sheet = resolve_sheet_name(source_name, source_content, sheet_name)
        return normalize_brazilian_decimals(read_excel_table(source_content, str(selected_sheet)))
    if source_extension == ".xls":
        return normalize_brazilian_decimals(pl.read_excel(BytesIO(source_content), sheet_name=sheet_name, infer_schema_length=500))
    if source_extension == ".json":
        return normalize_brazilian_decimals(read_json_table(source_content))
    raise HTTPException(415, "Formato não suportado. Envie ou indique CSV, XLSX, XLS, JSON, ou um .zip/.gz contendo um desses formatos.")

def profile_table(source_name: str, source_content: bytes, source_url: str | None, source_kind: str, sheet_name: str | None = None) -> dict[str, Any]:
    selected_sheet = resolve_sheet_name(source_name, source_content, sheet_name)
    source_table = read_table(source_name, source_content, selected_sheet)
    source_table = normalize_table_columns(source_table)
    if not source_table.height or not source_table.width:
        raise HTTPException(422, "Não encontramos linhas e colunas de dados nesta aba. Escolha uma aba que contenha uma tabela.")
    null_counts = {column_name: int(source_table[column_name].null_count()) for column_name in source_table.columns}
    source_profile = {
        "kind": source_kind,
        "file_name": source_name,
        "source_url": source_url,
        "rows": source_table.height,
        "columns": [{"name": column_name, "dtype": str(source_table[column_name].dtype), "null_count": null_counts[column_name]} for column_name in source_table.columns],
        "sample": source_table.head(20).to_dicts(),
        "suggestions": suggest_mapping(source_table),
        "indicator_recommendations": recommend_indicators(source_table),
    }
    if Path(source_name).suffix.lower() == ".xlsx":
        sheet_profiles = []
        for sheet in workbook_sheets(source_content):
            if sheet["name"] == selected_sheet:
                sheet_profiles.append({**sheet, "rows": source_table.height, "columns": source_table.width})
                continue
            sheet_profiles.append({**sheet, "rows": max(int(sheet["rows"]) - 1, 0), "rows_estimated": True})
        source_profile["sheets"] = sheet_profiles
        source_profile["selected_sheet"] = selected_sheet
        header_configuration = infer_excel_headers(source_content, str(selected_sheet))
        if header_configuration:
            source_profile["reading_notes"] = [f"O Decsys identificou {len(header_configuration['header_rows'])} linha(s) de cabeçalho e combinou os títulos antes de ler os dados."]
        source_profile["quality_warnings"] = excel_quality_warnings(source_content, str(selected_sheet), source_table)
    if len(source_content) > ai_assessment_size_limit:
        source_profile["agent_assessment"] = {"status": "skipped", "summary": f"A leitura estrutural foi concluída. O arquivo tem mais de {ai_assessment_size_limit // (1024 * 1024)} MB, então a avaliação por IA foi pulada para manter a resposta rápida."}
    else:
        source_profile["agent_assessment"] = ai_assessment.assess_source(source_profile)
    return source_profile

def excel_quality_warnings(source_content: bytes, sheet_name: str, source_table: pl.DataFrame) -> list[str]:
    quality_warnings: list[str] = []
    sparse_unlabeled_columns = []
    for column_name in source_table.columns:
        if not column_name.startswith("coluna_") or not source_table.height:
            continue
        populated_count = source_table.height - source_table[column_name].null_count()
        if populated_count and populated_count / source_table.height <= 0.1:
            sparse_unlabeled_columns.append(column_name)
    if sparse_unlabeled_columns:
        quality_warnings.append(f"Encontramos {len(sparse_unlabeled_columns)} coluna(s) sem título, com poucos valores. Confira se esses campos pertencem à tabela antes de mapear.")

    workbook = load_workbook(BytesIO(source_content), read_only=True, data_only=True)
    worksheet = workbook[sheet_name]
    title_text = " ".join(str(cell_value) for row_values in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row or 5, 5), values_only=True) for cell_value in row_values if has_excel_cell_value(cell_value))
    workbook.close()
    normalized_title_text = re.sub(r"[_\W]+", " ", title_text, flags=re.UNICODE)
    expected_count_match = re.search(r"\b(\d{1,4})\s+maiores?\b", normalized_title_text, re.IGNORECASE)
    if expected_count_match:
        expected_record_count = int(expected_count_match.group(1))
        if expected_record_count != source_table.height:
            quality_warnings.append(f"O título menciona {expected_record_count} registros, mas a tabela contém {source_table.height}. Confira se a planilha está completa antes de continuar.")
    return quality_warnings

def suggest_mapping(source_table: pl.DataFrame) -> dict[str, str]:
    suggestions: dict[str, str] = {}
    for column_name in source_table.columns:
        if is_municipality_code_column(column_name):
            suggestions.setdefault("municipality_code", column_name)
        elif "municipio" in column_name and not re.search(r"(?:maiores|menores|numero_de_municipios|participacao|posicao|total)", column_name):
            suggestions.setdefault("municipality_name", column_name)
        if column_name in {"ano", "year", "periodo", "ano_referencia"}:
            suggestions["reference_year"] = column_name
        if column_name in {"valor", "value", "indice", "percentual"}:
            suggestions["value"] = column_name
    return suggestions

def latest_measure_field(source_columns: list[str], measure_token: str) -> str | None:
    monthly_pattern = re.compile(r"(janeiro|fevereiro|marco|abril|maio|junho|julho|agosto|setembro|outubro|novembro|dezembro)_20[0-9]{2}.*" + re.escape(measure_token) + r"$")
    normalized_columns = [(column_name, normalized_sheet_title(normalize_column(column_name))) for column_name in source_columns]
    return next((column_name for column_name, normalized_column_name in reversed(normalized_columns) if monthly_pattern.search(normalized_column_name)), next((column_name for column_name, normalized_column_name in reversed(normalized_columns) if measure_token in normalized_column_name), None))

def recommend_indicators(source_table: pl.DataFrame) -> list[dict[str, str]]:
    source_columns = source_table.columns
    normalized_source_columns = [normalized_sheet_title(normalize_column(column_name)) for column_name in source_columns]
    source_columns_text = " ".join(normalized_source_columns)
    recommendations: list[dict[str, str]] = []
    if "codigo" in source_columns_text and "municipio" in source_columns_text and any(token in source_columns_text for token in ("admissoes", "desligamentos", "saldos", "estoque")):
        for measure_token, code, name, definition in (
            ("saldos", "caged_saldo_empregos", "Saldo de empregos formais", "Diferença entre admissões e desligamentos formais no período."),
            ("admissoes", "caged_admissoes", "Admissões formais", "Quantidade de admissões formais registradas no período."),
            ("desligamentos", "caged_desligamentos", "Desligamentos formais", "Quantidade de desligamentos formais registrados no período."),
            ("estoque", "caged_estoque_empregos", "Estoque de empregos formais", "Quantidade de vínculos formais de emprego no período."),
        ):
            value_field = latest_measure_field(source_columns, measure_token)
            if value_field:
                recommendations.append({"code": code, "name": name, "dimension": "trabalho", "definition": definition, "unit": "pessoas", "expected_frequency": "mensal", "value_field": value_field})
        return recommendations
    for column_name in source_columns:
        normalized_name = normalize_column(column_name)
        if "potencia" in normalized_name and normalized_name.endswith("kw"):
            recommendations.append({"code": "potencia_geracao_instalada", "name": "Potência de geração instalada", "dimension": "energia", "definition": "Potência declarada de empreendimentos de geração de energia.", "unit": "kW", "expected_frequency": "anual", "value_field": column_name})
            break
    return recommendations
