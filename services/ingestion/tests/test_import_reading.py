import json
import csv
import unittest
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from fastapi import HTTPException
from fastapi.testclient import TestClient
import polars as pl

from services.ingestion.app import main
from services.ingestion.app.main import import_tables, normalize_brazilian_decimals, normalize_table_columns, read_table, resolve_source_file_name, workbook_sheets


class ImportReadingTests(unittest.TestCase):
    def create_workbook(self) -> bytes:
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Sumário"
        summary_sheet.append(["Conteúdo da planilha"])
        source_sheet = workbook.create_sheet("Base municipal")
        source_sheet.append(["Base histórica de municípios"])
        source_sheet.append(["Código IBGE", "Município", "Ano", "Valor"])
        source_sheet.append(["3550308", "São Paulo", "2024", "1.234,5"])
        source_sheet.append(["3304557", "Rio de Janeiro", "2024", "2.345,6"])
        secondary_sheet = workbook.create_sheet("Outra tabela")
        secondary_sheet.append(["Código", "Indicador"])
        secondary_sheet.append(["A", "1,5"])
        workbook_buffer = BytesIO()
        workbook.save(workbook_buffer)
        return workbook_buffer.getvalue()

    def test_excel_finds_text_formatted_data_after_title_and_skips_summary(self):
        workbook_content = self.create_workbook()

        self.assertEqual([sheet["name"] for sheet in workbook_sheets(workbook_content)], ["Base municipal", "Outra tabela"])

        source_table = normalize_table_columns(read_table("municipios.xlsx", workbook_content, "Base municipal"))

        self.assertEqual(source_table.columns, ["codigo_ibge", "municipio", "ano", "valor"])
        self.assertEqual(source_table.height, 2)
        self.assertEqual(source_table.get_column("valor").to_list(), [1234.5, 2345.6])

    def test_excel_skips_report_titles_and_detects_numeric_text_rows(self):
        workbook = Workbook()
        source_sheet = workbook.active
        source_sheet.title = "Dados"
        source_sheet.append(["Tabela 1 - Maiores municípios", "__unnamed__1", "__unnamed__2", "__unnamed__3", "__unnamed__4"])
        source_sheet.append(["A preços correntes e participação acumulada"])
        source_sheet.append(["Segundo os municípios - 2021"])
        source_sheet.append([None, None, None, None, "(continua)"])
        source_sheet.append(["Municípios e respectivas_x000d_\nUnidades da Federação", "Posição ocupada", "Produto Interno Bruto_x000d_\n(1 000 R$)", "Participação (%)", "Participação_x000d_\nacumulada (%)"])
        source_sheet.append(["São Paulo (SP)", "1º", "828980607.731255293", "9.198485862", "9.198485862"])
        source_sheet.append(["Rio de Janeiro (RJ)", "2º", "359634752.586723745", "3.990557989", "13.189043851"])
        workbook_buffer = BytesIO()
        workbook.save(workbook_buffer)

        header_configuration = main.infer_excel_headers(workbook_buffer.getvalue(), "Dados")
        source_table = normalize_table_columns(read_table("tabelas.xlsx", workbook_buffer.getvalue(), "Dados"))

        self.assertEqual(header_configuration["header_rows"], [5])
        self.assertEqual(source_table.height, 2)
        self.assertEqual(source_table.columns, ["municipios_e_respectivas_unidades_da_federacao", "posicao_ocupada", "produto_interno_bruto_(1_000_r$)", "participacao_(%)", "participacao_acumulada_(%)"])
        self.assertEqual(source_table.get_column("produto_interno_bruto_(1_000_r$)").to_list(), [828980607.731255293, 359634752.586723745])

    def test_excel_discards_repeated_headers_inside_the_data(self):
        workbook = Workbook()
        source_sheet = workbook.active
        source_sheet.title = "Dados"
        source_sheet.append(["Município", "Ano", "Valor"])
        source_sheet.append(["Campinas", 2023, 10])
        source_sheet.append(["Município", "Ano", "Valor"])
        source_sheet.append(["Santos", 2024, 12])
        workbook_buffer = BytesIO()
        workbook.save(workbook_buffer)

        source_table = read_table("dados.xlsx", workbook_buffer.getvalue(), "Dados")

        self.assertEqual(source_table.height, 2)
        self.assertEqual(source_table.get_column("Município").to_list(), ["Campinas", "Santos"])

    def test_excel_combines_grouped_headers_and_keeps_sparse_columns(self):
        workbook = Workbook()
        source_sheet = workbook.active
        source_sheet.title = "Dados"
        source_sheet.append(["Relatório de indicadores"])
        source_sheet.append(["Período", None, "Indicador", None, None])
        source_sheet.append(["Município", "Ano", "Valor", None, None])
        source_sheet.append(["Campinas", 2023, 10, None, None])
        source_sheet.append(["Santos", 2024, 12, None, 99])
        workbook_buffer = BytesIO()
        workbook.save(workbook_buffer)

        source_table = normalize_table_columns(read_table("dados.xlsx", workbook_buffer.getvalue(), "Dados"))

        self.assertEqual(source_table.height, 2)
        self.assertEqual(source_table.columns, ["periodo_municipio", "periodo_ano", "indicador_valor", "coluna_4"])
        self.assertEqual(source_table.get_column("coluna_4").to_list(), [None, 99])

    def test_excel_quality_warnings_flag_incomplete_titles_and_sparse_columns(self):
        workbook = Workbook()
        source_sheet = workbook.active
        source_sheet.title = "Dados"
        source_sheet.append(["Tabela_1_-_Posição_ocupada_pelos_100_maiores_municípios"])
        source_sheet.append(["Município", "Ano", "Valor", None, None])
        for record_number in range(11):
            sparse_value = 99 if record_number == 10 else None
            source_sheet.append([f"Município {record_number + 1}", 2020 + record_number, record_number + 1, None, sparse_value])
        workbook_buffer = BytesIO()
        workbook.save(workbook_buffer)
        workbook_content = workbook_buffer.getvalue()
        source_table = normalize_table_columns(read_table("dados.xlsx", workbook_content, "Dados"))

        quality_warnings = main.excel_quality_warnings(workbook_content, "Dados", source_table)

        self.assertEqual(source_table.height, 11)
        self.assertEqual(len(quality_warnings), 2)
        self.assertIn("coluna(s) sem título", quality_warnings[0])
        self.assertIn("100", quality_warnings[1])
        self.assertIn("11", quality_warnings[1])

    def test_all_workbook_sheets_keep_their_own_rows_and_names(self):
        workbook_content = self.create_workbook()

        imported_tables = import_tables("municipios.xlsx", workbook_content, "Base municipal", True)

        self.assertEqual([sheet_name for sheet_name, _ in imported_tables], ["Base municipal", "Outra tabela"])
        self.assertEqual(imported_tables[0][1].height, 2)

    def test_automatic_sheet_choice_uses_readable_data_volume(self):
        workbook = Workbook()
        source_sheet = workbook.active
        source_sheet.title = "Base principal"
        source_sheet.append(["Código", "Município", "Ano", "Valor"])
        source_sheet.append(["3550308", "São Paulo", "2024", "1"])
        source_sheet.append(["3304557", "Rio", "2024", "2"])
        notes_sheet = workbook.create_sheet("Notas longas")
        notes_sheet.append(["Código", "Valor"])
        notes_sheet.append(["A", "1"])
        for note_number in range(80):
            notes_sheet.append([f"Nota {note_number}"])
        workbook_buffer = BytesIO()
        workbook.save(workbook_buffer)

        imported_tables = import_tables("municipios.xlsx", workbook_buffer.getvalue(), None, False)

        self.assertEqual(imported_tables[0][0], "Base principal")
        self.assertEqual(imported_tables[0][1].height, 2)

    def test_csv_supports_brazilian_separator_decimal_and_windows_encoding(self):
        csv_content = "Município;Ano;Valor\r\nCampinas;2024;1.234,5\r\nSantos;2024;2.345,6\r\n".encode("cp1252")

        source_table = normalize_table_columns(read_table("municipios.csv", csv_content))

        self.assertEqual(source_table.columns, ["municipio", "ano", "valor"])
        self.assertEqual(source_table.get_column("valor").to_list(), [1234.5, 2345.6])

    def test_brazilian_decimal_normalization_keeps_valid_rows_when_some_cells_are_invalid(self):
        source_table = normalize_brazilian_decimals(read_table("dados.csv", b"valor;categoria\n1.25;A\n2,5;B\nnao informado;C\n"))

        self.assertEqual(source_table.get_column("valor").to_list(), [1.25, 2.5, None])

    def test_ibge_codes_stay_seven_digit_strings_when_excel_infers_numbers(self):
        source_table = pl.DataFrame({"Código IBGE": [3550308.0, None, 120001.0]})

        normalized_table = normalize_table_columns(source_table)

        self.assertEqual(normalized_table.get_column("codigo_ibge").to_list(), ["3550308", None, "0120001"])

    def test_json_rows_and_sidra_label_rows_are_readable(self):
        plain_json = json.dumps([{"municipio": "Campinas", "ano": 2024, "valor": 12.5}]).encode()
        sidra_json = json.dumps([
            {"D1N": "Município (Código)", "V": "Variável (Código)"},
            {"D1N": "Campinas (3509502)", "V": "12,5"},
        ], ensure_ascii=False).encode()

        plain_table = read_table("dados.json", plain_json)
        sidra_table = read_table("sidra.json", sidra_json)

        self.assertEqual(plain_table.height, 1)
        self.assertEqual(sidra_table.columns, ["Município (Código)", "Variável (Código)"])
        self.assertEqual(sidra_table.height, 1)

    def test_zip_selects_a_supported_data_file(self):
        archive_buffer = BytesIO()
        with ZipFile(archive_buffer, "w", ZIP_DEFLATED) as source_archive:
            source_archive.writestr("readme.txt", "not a data table")
            source_archive.writestr("dados.csv", "municipio;ano;valor\nCampinas;2024;1,5\n")

        source_table = read_table("dados.zip", archive_buffer.getvalue())

        self.assertEqual(source_table.height, 1)
        self.assertEqual(source_table.get_column("valor").to_list(), [1.5])

    def test_link_download_names_can_come_from_headers_or_excel_mime_types(self):
        header_name = resolve_source_file_name("https://example.test/download?id=1", "application/octet-stream", "attachment; filename*=UTF-8''dados%20municipais.xlsx")
        mime_name = resolve_source_file_name("https://example.test/download?id=1", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        self.assertEqual(header_name, "dados municipais.xlsx")
        self.assertEqual(mime_name, "download.xlsx")

    @patch("services.ingestion.app.main.assess_source", return_value={"status": "skipped", "summary": "Teste local"})
    def test_profile_endpoint_returns_the_parsed_csv_structure(self, assess_source):
        client = TestClient(main.app)

        response = client.post(
            "/profile",
            files={"file": ("municipios.csv", "codigo_ibge;ano;valor\n3550308;2024;1.234,5\n", "text/csv")},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["rows"], 1)
        self.assertEqual([column["name"] for column in response.json()["columns"]], ["codigo_ibge", "ano", "valor"])
        self.assertEqual(response.json()["sample"][0]["valor"], 1234.5)

    def test_keepalive_endpoint_rejects_requests_without_the_shared_secret(self):
        with patch.dict("os.environ", {"KEEPALIVE_SECRET": "local-test-secret"}):
            response = TestClient(main.app).get("/keepalive")

        self.assertEqual(response.status_code, 401)

    def test_keepalive_endpoint_runs_a_read_only_database_query(self):
        with patch.dict("os.environ", {"KEEPALIVE_SECRET": "local-test-secret"}):
            with patch("services.ingestion.app.main.supabase_headers", return_value={}), patch("services.ingestion.app.main.supabase_url", return_value="https://supabase.test/rest/v1/imports"), patch("services.ingestion.app.main.httpx.get", return_value=self.successful_response([])) as database_query:
                response = TestClient(main.app).get("/keepalive", headers={"Authorization": "Bearer local-test-secret"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "ok", "database": "reachable"})
        self.assertEqual(database_query.call_args.kwargs["params"], {"select": "id", "limit": "1"})

    @patch("services.ingestion.app.main.assess_source", return_value={"status": "skipped", "summary": "Teste local"})
    @patch("services.ingestion.app.main.supabase_headers", return_value={})
    @patch("services.ingestion.app.main.supabase_url", side_effect=lambda path: f"https://supabase.test{path}")
    @patch("services.ingestion.app.main.httpx.post")
    @patch("services.ingestion.app.main.httpx.patch")
    def test_draft_persists_all_readable_sheets_and_keeps_sheet_row_numbers(self, patch_request, post_request, supabase_url, supabase_headers, assess_source):
        post_request.side_effect = [
            self.successful_response([{"id": "source-id"}]),
            self.successful_response([{"id": "import-id"}]),
            self.successful_response([]),
            self.successful_response([]),
            self.successful_response([]),
        ]
        patch_request.return_value = self.successful_response([])

        result = main.persist_import("municipios.xlsx", self.create_workbook(), None, None, "Municípios", None, "Base municipal", True)

        self.assertEqual(result["total_rows"], 3)
        self.assertEqual(result["imported_sheets"], ["Base municipal", "Outra tabela"])
        sheet_records = post_request.call_args_list[2].kwargs["json"]
        self.assertEqual([(sheet["sheet_name"], sheet["selected_for_treatment"]) for sheet in sheet_records], [("Base municipal", True), ("Outra tabela", False)])
        staged_row_records = post_request.call_args_list[4].kwargs["json"]
        self.assertEqual([(row["sheet_name"], row["row_number"]) for row in staged_row_records], [("Base municipal", 1), ("Base municipal", 2), ("Outra tabela", 1)])

    def test_duplicate_and_empty_column_names_are_made_unique(self):
        source_table = read_table("dados.csv", b";Valor;Valor\nA;1;2\n")

        normalized_table = normalize_table_columns(source_table)

        self.assertEqual(normalized_table.columns, ["coluna_1", "valor", "valor_2"])

    def test_empty_and_corrupt_files_return_a_reading_error(self):
        for source_content in (b"", b"not an xlsx"):
            with self.subTest(source_content=source_content):
                with self.assertRaises(HTTPException) as caught_error:
                    read_table("dados.xlsx", source_content)
                self.assertEqual(caught_error.exception.status_code, 422)

    @patch("services.ingestion.app.main.fetch_import_rows")
    def test_xlsx_export_keeps_each_imported_sheet_separate(self, fetch_rows):
        fetch_rows.return_value = [
            {"sheet_name": "Municípios", "row_number": 1, "raw_row": {"codigo": "3550308", "valor": 12.5}},
            {"sheet_name": "Indicadores", "row_number": 1, "raw_row": {"indicador": "A", "valor": 8}},
        ]

        response = main.export_import_xlsx("import-id")
        workbook = load_workbook(BytesIO(response.body), data_only=True)

        self.assertEqual(workbook.sheetnames, ["Municípios", "Indicadores"])
        self.assertEqual(list(workbook["Municípios"].values), [("codigo", "valor"), ("3550308", 12.5)])
        self.assertEqual(list(workbook["Indicadores"].values), [("indicador", "valor"), ("A", 8)])

    @patch("services.ingestion.app.main.fetch_import_rows")
    def test_csv_export_marks_sheet_and_serializes_nested_fields(self, fetch_rows):
        fetch_rows.return_value = [
            {"sheet_name": "Uma", "row_number": 1, "raw_row": {"aba": "valor de origem", "detalhes": {"fonte": "IBGE"}}},
            {"sheet_name": "Duas", "row_number": 1, "raw_row": {"aba": "outro valor", "detalhes": [1, 2]}},
        ]

        response = main.export_import_csv("import-id")
        exported_rows = list(csv.reader(response.body.decode("utf-8-sig").splitlines(), delimiter=";"))

        self.assertEqual(exported_rows[0], ["aba_2", "aba", "detalhes"])
        self.assertEqual(exported_rows[1], ["Uma", "valor de origem", '{"fonte":"IBGE"}'])
        self.assertEqual(exported_rows[2], ["Duas", "outro valor", "[1,2]"])

    @patch("services.ingestion.app.main.httpx.get")
    @patch("services.ingestion.app.main.supabase_headers", return_value={})
    @patch("services.ingestion.app.main.supabase_url", side_effect=lambda path: f"https://supabase.test{path}")
    def test_export_reads_every_supabase_page(self, supabase_url, supabase_headers, get_request):
        first_page = [{"sheet_name": "Dados", "row_number": row_number, "raw_row": {"valor": row_number}} for row_number in range(1, 1001)]
        second_page = [{"sheet_name": "Dados", "row_number": 1001, "raw_row": {"valor": 1001}}]
        get_request.side_effect = [self.successful_response(first_page), self.successful_response(second_page)]

        imported_rows = main.fetch_import_rows("import-id")

        self.assertEqual(len(imported_rows), 1001)
        self.assertEqual(get_request.call_args_list[0].kwargs["params"]["offset"], "0")
        self.assertEqual(get_request.call_args_list[1].kwargs["params"]["offset"], "1000")

    @staticmethod
    def successful_response(response_json):
        class Response:
            is_success = True

            def json(self):
                return response_json

        return Response()


if __name__ == "__main__":
    unittest.main()
